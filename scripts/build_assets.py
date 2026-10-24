"""Regenerates excel/AutomatedInvoiceProcessingAgentOutput.xlsx and sample-data/*.pdf.
The Excel columns come from OUTPUT_COLUMNS in the MCP server so they never drift.
Usage: python scripts/build_assets.py   (needs: openpyxl reportlab)"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src/mcp-server"))
from invoiceiq.models import COLUMN_DESCRIPTIONS, OUTPUT_COLUMNS  # noqa: E402

# ---------------------------------------------------------------- Excel register
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADER_ROW = 4
last_col = get_column_letter(len(OUTPUT_COLUMNS))
wb = Workbook()
ws = wb.active
ws.title = "Invoice Output"

ws["A1"] = "Automated Invoice Processing Agent Output"
ws["A1"].font = Font(name="Cambria", size=16, bold=True, color="FFFFFF")
ws["A2"] = "Structured output table for invoice information extracted by the Copilot Studio agent."
ws["A2"].font = Font(name="Cambria", italic=True, color="1F3864")
for row, fill in ((1, "1F4E79"), (2, "DDEBF7")):
    for c in range(1, len(OUTPUT_COLUMNS) + 1):
        ws.cell(row, c).fill = PatternFill("solid", fgColor=fill)
ws.row_dimensions[1].height = 26

for i, col in enumerate(OUTPUT_COLUMNS, 1):
    cell = ws.cell(HEADER_ROW, i, col)
    cell.comment = Comment(COLUMN_DESCRIPTIONS[col], "InvoiceIQ")
    cell.alignment = Alignment(horizontal="center", vertical="center")
# one empty data row - an Excel table needs at least one body row; the flow appends below it
tbl = Table(displayName="InvoiceOutput", ref=f"A{HEADER_ROW}:{last_col}{HEADER_ROW + 1}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
ws.row_dimensions[HEADER_ROW].height = 30

widths = {"VendorName": 30, "InvoiceNumber": 22, "InvoiceDate": 14, "CustomerName": 30, "TotalTax": 12,
          "InvoiceTotal": 14, "Currency": 10, "ConfidenceScore": 16, "DocumentURL": 70}
for i, col in enumerate(OUTPUT_COLUMNS, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)
ws.freeze_panes = f"A{HEADER_ROW + 1}"

notes = wb.create_sheet("ReadMe")
for line in [
    "Written by the Copilot Studio agent flow 'Save Automated Invoice Output' (Excel Online: Add a row into a table).",
    "Table name: InvoiceOutput  |  Sheet: Invoice Output  - do not rename the table or its columns.",
    "Missing values are written as NA.  DocumentURL is stored without the SAS token.",
    "Microsoft Fabric Dataflow Gen2 reads this table (see fabric/ in the repository).",
]:
    notes.append([line])
notes.column_dimensions["A"].width = 120
(ROOT / "excel").mkdir(exist_ok=True)
wb.save(ROOT / "excel/AutomatedInvoiceProcessingAgentOutput.xlsx")

# ---------------------------------------------------------------- sample invoices
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def invoice_pdf(path, vendor, address, inv_no, inv_date, due, bill_to, items, tax_label, tax_rate, cur):
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 20); c.drawString(50, h - 60, vendor)
    c.setFont("Helvetica", 10); c.drawString(50, h - 78, address)
    c.setFont("Helvetica-Bold", 16); c.drawString(410, h - 60, "INVOICE")
    c.setFont("Helvetica", 10)
    for i, (k, v) in enumerate([("Invoice No:", inv_no), ("Invoice Date:", inv_date), ("Due Date:", due)]):
        c.drawString(400, h - 85 - i * 14, k); c.drawString(475, h - 85 - i * 14, v)
    c.drawString(50, h - 130, "Bill To:"); c.drawString(50, h - 144, bill_to)
    y = h - 200
    c.setFont("Helvetica-Bold", 10)
    for x, t in [(50, "Description"), (320, "Qty"), (380, "Unit Price"), (480, "Amount")]:
        c.drawString(x, y, t)
    c.line(50, y - 4, 545, y - 4); c.setFont("Helvetica", 10)
    sub = 0
    for desc, qty, price in items:
        y -= 18; amt = qty * price; sub += amt
        c.drawString(50, y, desc); c.drawRightString(340, y, str(qty))
        c.drawRightString(440, y, f"{price:,.2f}"); c.drawRightString(545, y, f"{amt:,.2f}")
    tax = round(sub * tax_rate, 2)
    y -= 30
    rows = [("Subtotal", sub)] + ([(tax_label, tax)] if tax_rate else []) + [(f"Total ({cur})", sub + tax)]
    for label, val in rows:
        c.drawString(380, y, label); c.drawRightString(545, y, f"{cur} {val:,.2f}"); y -= 16
    c.save()


out = ROOT / "sample-data"; out.mkdir(exist_ok=True)
invoice_pdf(out / "01_Contoso_Office_Supplies_INR.pdf", "Contoso Supplies Pvt Ltd",
            "12 Mount Road, Chennai 600002, India   GSTIN: 33AABCC1234F1Z5", "CS-2026-1042", "05-10-2026",
            "04-11-2026", "Fabrikam Industries Ltd, Bengaluru",
            [("A4 Copier Paper (box)", 20, 450.0), ("Laser Toner Cartridge", 3, 3200.0), ("Desk Organiser", 5, 380.0)],
            "GST 18%", 0.18, "INR")
invoice_pdf(out / "02_Northwind_Cloud_Services_USD.pdf", "Northwind Cloud Services LLC",
            "500 Pine Street, Seattle, WA 98101, USA", "NW-88213", "Oct 12, 2026", "Nov 11, 2026",
            "Adventure Works Cycles India - Accounts Payable",
            [("Managed Kubernetes - October", 1, 2400.0), ("Premium Support", 1, 650.0)], "", 0.0, "USD")
invoice_pdf(out / "03_Tailspin_Consulting_GBP.pdf", "Tailspin Consulting Ltd",
            "221 King's Road, London SW3 5EL, UK   VAT: GB123456789", "TSC-2026-0457", "15/10/2026", "14/11/2026",
            "Fabrikam Manufacturing Ltd",
            [("Architecture workshop (days)", 3, 950.0), ("Travel & expenses", 1, 310.0)], "VAT 20%", 0.20, "GBP")
print("Columns:", OUTPUT_COLUMNS)
